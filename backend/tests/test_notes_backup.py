"""The daily plain-file export of notes, lists and tables (notes/backup.py) and its readiness
check `notes_export` (core/health.py)."""

from __future__ import annotations

import csv
import io
import json
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook

from mailr.config import get_settings
from mailr.core import backup as mail_backup
from mailr.core import health
from mailr.notes import backup

# These tests never touch the mail store: database-only users (conftest.make_db_user).
pytestmark = pytest.mark.db_user

API = "/api/v1"


@pytest.fixture
def dirs(tmp_path, monkeypatch):
    s = get_settings()
    monkeypatch.setattr(s, "data_dir", str(tmp_path / "data"))
    monkeypatch.setattr(s, "backup_target", "")
    return tmp_path


async def _content(client):
    r = await client.post(f"{API}/notes", json={"kind": "checklist", "title": "Groceries", "lines": "milk ×2\neggs"})
    assert r.status_code == 200, r.text
    r = await client.post(f"{API}/notes", json={"kind": "text", "title": "Wifi", "body": "The password is on the fridge."})
    assert r.status_code == 200, r.text
    r = await client.post(f"{API}/tables", json={"name": "Trip costs", "columns": [
        {"name": "Item", "type": "text"}, {"name": "Cost", "type": "number"}]})
    assert r.status_code == 200, r.text
    t = r.json()["table"]
    cols = {c["name"]: c["id"] for c in t["columns"]}
    r = await client.post(f"{API}/tables/{t['id']}/rows", json={"cells": {cols["Item"]: "Tent", cols["Cost"]: 120}})
    assert r.status_code == 200, r.text
    return t


async def test_export_writes_markdown_csv_xlsx_and_json(app, user, client, dirs):
    await _content(client)
    out = await backup.export_user(user.id)
    assert out["snapshot"] and out["counts"] == {"notes": 1, "lists": 1, "tables": 1}
    snap = backup.export_dir(user.id) / out["snapshot"]
    md = (snap / "notes" / "Groceries.md").read_text()
    assert md.startswith("# Groceries") and "- [ ] milk ×2" in md and "- [ ] eggs" in md
    assert "fridge" in (snap / "notes" / "Wifi.md").read_text()
    rows = list(csv.reader(io.StringIO((snap / "tables" / "Trip costs.csv").read_text().lstrip("﻿"))))
    assert rows[0][:2] == ["Item", "Cost"] and rows[1][0] == "Tent"
    wb = load_workbook(snap / "tables.xlsx")
    assert wb.sheetnames == ["Trip costs"]
    full = json.loads((snap / "tables.json").read_text())["tables"][0]
    assert full["name"] == "Trip costs" and len(full["rows"]) == 1 and full["columns"]
    notes = json.loads((snap / "notes.json").read_text())["notes"]
    assert {n["title"] for n in notes} == {"Groceries", "Wifi"}
    manifest = json.loads((snap / "manifest.json").read_text())
    assert "notes/Groceries.md" in manifest["files"] and manifest["format"] == backup.FORMAT
    assert mail_backup.last_run(backup.export_dir(user.id)) is not None

    # Nothing changed: no new snapshot, the run is still recorded.
    again = await backup.export_user(user.id)
    assert again["snapshot"] is None


async def test_the_copy_on_the_target(app, user, client, dirs, monkeypatch):
    await _content(client)
    tgt = dirs / "nas"
    tgt.mkdir()
    (tgt / mail_backup.TARGET_MARKER).write_text("")
    monkeypatch.setattr(get_settings(), "backup_target", str(tgt))
    out = await backup.export_user(user.id)
    assert out["copied"] == out["snapshot"]
    assert (tgt / "notes" / f"user-{user.id}" / out["snapshot"] / "notes" / "Groceries.md").is_file()
    assert mail_backup.last_run(backup.target_dir(user.id)) is not None


def test_the_readiness_rule():
    now = datetime.now(UTC)
    old = now - timedelta(hours=30)
    assert health.evaluate_notes_runs([], None, now) == (True, "no notes or tables yet")
    assert health.evaluate_notes_runs([now, None], None, now) == (False, "1 of 2 notes exports not run in 26 h")
    assert health.evaluate_notes_runs([now], [old], now) == (False, "1 of 1 notes export not copied in 26 h")
    assert health.evaluate_notes_runs([now, now], [now, now], now) == (True, "2 notes exports in the last 26 h and copied")
    assert "notes_export" in health.CHECKS and "notes_export" in health.BACKUP_CHECKS


async def test_the_check_follows_the_runs(app, user, client, dirs, monkeypatch):
    await _content(client)

    async def only_this_user():
        return [user.id]

    monkeypatch.setattr(backup, "owners_with_data", only_this_user)
    # New notes wait for their first export (core/health.py _stale_waiting); a day later they fail.
    assert await health.check_notes_export() == (True, "0 of 1 notes export in the last 26 h, 1 new (first export pending)")
    from sqlalchemy import text

    from mailr.db import session_scope

    async with session_scope() as db:
        for t in ("notes", "tables"):
            await db.execute(text(f"UPDATE {t} SET created_at = now() - interval '30 hours' WHERE created_by = :u"),
                             {"u": user.id})
    assert await health.check_notes_export() == (False, "1 of 1 notes export not run in 26 h")
    await backup.export_user(user.id)
    assert await health.check_notes_export() == (True, "1 notes export in the last 26 h")
    res = await health.run_checks(("notes_export",))
    assert res["checks"]["notes_export"]["ok"] and res["checks"]["notes_export"]["required"]
