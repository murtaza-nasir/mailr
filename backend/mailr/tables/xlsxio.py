""".xlsx export.

One sheet per table, named after it. Cells keep their types: numbers and amounts as numbers (an
amount with its currency's number format), dates and date-times as dates, checkboxes as TRUE or
FALSE, multi-select options joined with "; ", links as their titles. The header row is bold and
frozen, and each column is about as wide as its longest value. A cell that holds text the column's
type could not read (the `raw` of an import) is written as that text.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from mailr.tables import types as ty

MIN_WIDTH, MAX_WIDTH = 8, 60
# Excel's number formats for the currencies a symbol stands for; any other code is written after
# the amount ("1,234.50 SEK").
SYMBOL = {"USD": "$", "EUR": "€", "GBP": "£", "JPY": "¥", "INR": "₹", "CAD": "C$", "AUD": "A$", "KRW": "₩",
          "BRL": "R$", "ILS": "₪", "TRY": "₺"}
LINK_WORD = {"row": "row", "conversation": "email", "event": "event", "task": "task", "note": "note"}


def sheet_title(name: str, taken: set[str]) -> str:
    """A valid, unique sheet name: at most 31 characters, none of []:*?/\\."""
    base = re.sub(r"[\[\]:*?/\\]", " ", name or "").strip(" '") or "Table"
    base = base[:31]
    title, n = base, 2
    while title.lower() in taken:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    taken.add(title.lower())
    return title


def number_format(col: ty.Col) -> str | None:
    cfg = col.config or {}
    places = cfg.get("decimals")
    if col.type == "currency":
        dec = "0" * int(2 if places is None else places)
        num = "#,##0" + (f".{dec}" if dec else "")
        code = cfg.get("currency") or ""
        sym = SYMBOL.get(code)
        return f'"{sym}"{num}' if sym else (f'{num} "{code}"' if code else num)
    fmt = cfg.get("format") or "plain"
    dec = "0" * int(places) if places is not None else ""
    if fmt == "percent":
        return "0" + (f".{dec}" if dec else "") + "%"
    if fmt == "thousands":
        return "#,##0" + (f".{dec}" if dec else "")
    return ("0." + dec) if dec else None


def cell_value(col: ty.Col, v: Any, raw: str | None, titles: dict[str, str]) -> Any:
    """The typed value of one cell; None for an empty one."""
    if ty.is_empty(v):
        if col.type == "checkbox" and not raw:
            return False
        return raw or None
    t = col.type
    if t in ("number", "currency"):
        d = Decimal(str(v))
        return int(d) if d == d.to_integral_value() and abs(d) < 10**15 else float(d)
    if t == "date":
        return date.fromisoformat(str(v))
    if t == "datetime":
        # The wall time the grid shows; Excel has no time zones.
        at = datetime.fromisoformat(str(v.get("at", "")))
        return at.replace(tzinfo=None)
    if t == "checkbox":
        return bool(v)
    if t == "link":
        return "; ".join(titles.get(ty.ref_key(r), "") or LINK_WORD.get(r.get("kind"), "link") for r in v)
    return ty.plain(col, v, titles=titles)


def _width(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, bool):
        return 5
    if isinstance(value, datetime):
        return 16
    if isinstance(value, date):
        return 10
    if isinstance(value, (int, float)):
        return len(f"{value:,.2f}")
    return max((len(x) for x in str(value).splitlines()), default=0)


def add_sheet(wb: Workbook, name: str, columns: list[ty.Col], rows: list[Any], titles: dict[str, str],
              taken: set[str]) -> None:
    ws = wb.create_sheet(sheet_title(name, taken))
    bold = Font(bold=True)
    widths = []
    for j, c in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=j, value=c.name)
        cell.data_type = "s"  # text, never a formula
        cell.font = bold
        widths.append(len(c.name))
    fmts = [number_format(c) if c.type in ("number", "currency") else None for c in columns]
    for i, r in enumerate(rows, start=2):
        for j, c in enumerate(columns, start=1):
            v = cell_value(c, (r.cells or {}).get(c.col_id), (r.raw or {}).get(c.col_id), titles)
            if v is None:
                continue
            cell = ws.cell(row=i, column=j, value=v)
            if isinstance(v, str):
                # Text stays text: "=HYPERLINK(…)" or "=WEBSERVICE(…)" from a sharee or from mail is
                # never a formula in the spreadsheet (openpyxl makes one of any text starting "=").
                cell.data_type = "s"
            if isinstance(v, (int, float)) and not isinstance(v, bool) and fmts[j - 1]:
                cell.number_format = fmts[j - 1]
            elif isinstance(v, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(v, date):
                cell.number_format = "yyyy-mm-dd"
            widths[j - 1] = max(widths[j - 1], _width(v))
    ws.freeze_panes = "A2"
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(MIN_WIDTH, min(MAX_WIDTH, w + 2))


def workbook(sheets: list[tuple[str, list[ty.Col], list[Any], dict[str, str]]]) -> bytes:
    """An .xlsx file with one sheet per (name, columns, rows, link titles)."""
    wb = Workbook()
    wb.remove(wb.active)
    taken: set[str] = set()
    for name, cols, rows, titles in sheets:
        add_sheet(wb, name, cols, rows, titles, taken)
    if not wb.worksheets:
        wb.create_sheet("Table")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
